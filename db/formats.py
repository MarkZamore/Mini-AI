"""Native structure extraction. No network calls or document rewriting."""
import hashlib, io, math, os, posixpath, re, subprocess, zipfile, tempfile, uuid
from collections import defaultdict
from pathlib import Path
from lxml import etree as E
from common import artifact, native, relative, sha256, CREATE_NO_WINDOW

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'a':'http://schemas.openxmlformats.org/drawingml/2006/main','s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing','mc':'http://schemas.openxmlformats.org/markup-compatibility/2006'}
W='{'+NS['w']+'}'
_OFFICE_STATE=None

def stop_office():
    global _OFFICE_STATE
    if _OFFICE_STATE is None:return
    import psutil
    try:
        parent=psutil.Process(_OFFICE_STATE['listener'].pid)
        children=parent.children(recursive=True)
        for child in reversed(children):
            try:child.terminate()
            except psutil.Error:pass
        parent.terminate()
        psutil.wait_procs(children+[parent],timeout=10)
    except psutil.Error:pass
    _OFFICE_STATE=None

def warning(code,message): return {'code':code,'message':message}
def quality(method='native',confidence=None,warnings=None,review=False):
    return {'method':method,'confidence':confidence,'needs_review':review,'verification':'machine_extracted','warnings':warnings or []}

class Context:
    def __init__(self,document,root): self.doc=document; self.root=Path(root); self.serial=0
    def loc(self,artifact_id='original',kind='artifact',**kwargs):
        return {'artifact_id':artifact_id,'source_sha256':self.doc['source_version_sha256'],'kind':kind,**kwargs}
    def block(self,kind,text,refs,parent=None,role='body',payload=None,q=None):
        address=hashlib.sha256(str(refs).encode()).hexdigest()[:16]
        bid=f"{self.doc['document_id']}:{self.doc['source_version_sha256']}:{address}:{self.serial}"
        b={'id':bid,'order':self.serial,'parent_id':parent,'kind':kind,'role':role,'text':text,'payload':payload or {},'source_refs':refs,'quality':q or quality()}
        self.serial+=1; self.doc['blocks'].append(b); return b
    def warn(self,code,message): self.doc['extraction']['warnings'].append(warning(code,message))
    def save_bytes(self,data,ext,derived_from,kind,refs=None):
        digest=hashlib.sha256(data).hexdigest(); ext=re.sub('[^a-zA-Z0-9]','',ext) or 'bin'
        path=self.root/'artifacts'/digest[:2]/(digest+'.'+ext)
        path.parent.mkdir(parents=True,exist_ok=True)
        if not path.exists():
            fd,tmp=tempfile.mkstemp(suffix='.part',dir=path.parent)
            with os.fdopen(fd,'wb') as f: f.write(data); f.flush(); os.fsync(f.fileno())
            os.replace(tmp,path)
        a=artifact(self.root,path,derived_from,kind)
        if not any(x['id']==a['id'] for x in self.doc['artifacts']): self.doc['artifacts'].append(a)
        if refs:
            asset={'id':'asset:'+digest,'artifact_id':a['id'],'role':kind,'source_refs':refs}
            if asset not in self.doc['assets']: self.doc['assets'].append(asset)
        return a,path

def parse_xml(data): return E.fromstring(data,parser=E.XMLParser(resolve_entities=False,no_network=True,huge_tree=True))
def xml_ref(ctx,aid,part,node): return ctx.loc(aid,'ooxml_node',part=part,xpath=node.getroottree().getpath(node))
def selected(node):
    for ac in list(node.iter('{'+NS['mc']+'}AlternateContent')):
        choices=[c for c in ac if E.QName(c).localname=='Choice']
        fallback=next((c for c in ac if E.QName(c).localname=='Fallback'),None)
        branch=choices[0] if choices else fallback
        if branch is not None:
            for child in list(ac):
                if child is not branch: ac.remove(child)
    return node

def text_of(node):
    out=[]
    def visit(n):
        tag=E.QName(n).localname
        if tag in {'txbxContent','tbl'} and n is not node: return
        if n.tag in {W+'t',W+'delText'}: out.append(n.text or '')
        elif n.tag==W+'tab': out.append('\t')
        elif n.tag in {W+'br',W+'cr'}: out.append('\n')
        else:
            for child in n: visit(child)
    visit(node); return ''.join(out)

def relationships(z,part):
    name=posixpath.join(posixpath.dirname(part),'_rels',posixpath.basename(part)+'.rels')
    if name not in z.namelist(): return {}
    return {n.get('Id'):{'target':n.get('Target'),'external':n.get('TargetMode')=='External','type':n.get('Type','')} for n in parse_xml(z.read(name))}

def extract_docx(ctx,path,aid='original',parent=None,depth=0):
    if depth>4: ctx.warn('embedding_depth','Embedded document recursion limit'); return
    with zipfile.ZipFile(native(path)) as z:
        styles={}; numbering={}; nums={}; counters=defaultdict(lambda:[0]*9)
        if 'word/styles.xml' in z.namelist():
            for st in parse_xml(z.read('word/styles.xml')).findall('w:style',NS):
                name=st.find('w:name',NS); outline=st.find('w:pPr/w:outlineLvl',NS)
                styles[st.get(W+'styleId')]={'name':name.get(W+'val','') if name is not None else '', 'level':int(outline.get(W+'val'))+1 if outline is not None else None}
        if 'word/numbering.xml' in z.namelist():
            nt=parse_xml(z.read('word/numbering.xml'))
            for abstract in nt.findall('w:abstractNum',NS):
                levels={}
                for level in abstract.findall('w:lvl',NS):
                    levels[int(level.get(W+'ilvl'))]={E.QName(c).localname:c.get(W+'val') for c in level if c.get(W+'val') is not None}
                numbering[abstract.get(W+'abstractNumId')]=levels
            for num in nt.findall('w:num',NS):
                abstract=num.find('w:abstractNumId',NS)
                if abstract is not None: nums[num.get(W+'numId')]=numbering.get(abstract.get(W+'val'),{})
        parts=['word/document.xml']+[n for n in z.namelist() if re.fullmatch(r'word/(header\d+|footer\d+|footnotes|endnotes|comments)\.xml',n)]
        for part in parts:
            root=selected(parse_xml(z.read(part))); rels=relationships(z,part)
            role='body' if part=='word/document.xml' else ('header' if '/header' in part else 'footer' if '/footer' in part else 'note')
            def media(node,owner):
                for image in node.xpath('.//a:blip|.//*[local-name()="imagedata"]',namespaces=NS):
                    rid=image.get('{'+NS['r']+'}embed') or image.get('{'+NS['r']+'}id')
                    rel=rels.get(rid)
                    if not rel or rel['external']: continue
                    target=posixpath.normpath(posixpath.join(posixpath.dirname(part),rel['target']))
                    if target not in z.namelist(): ctx.warn('missing_media',target); continue
                    ref=xml_ref(ctx,aid,part,image)
                    a,p=ctx.save_bytes(z.read(target),Path(target).suffix.lstrip('.'),aid,'image',[ref])
                    ctx.block('image','',[ref],owner,role,{'asset_id':'asset:'+a['sha256'],'artifact_id':a['id'],'embedded_part':target})
                for obj in node.xpath('.//*[local-name()="OLEObject"]'):
                    rid=obj.get('{'+NS['r']+'}id'); rel=rels.get(rid)
                    if not rel or rel['external']: continue
                    target=posixpath.normpath(posixpath.join(posixpath.dirname(part),rel['target']))
                    if target not in z.namelist(): continue
                    ref=xml_ref(ctx,aid,part,obj); data=z.read(target)
                    a,p=ctx.save_bytes(data,Path(target).suffix.lstrip('.'),aid,'attachment',[ref])
                    b=ctx.block('attachment','',[ref],owner,role,{'artifact_id':a['id'],'embedded_part':target})
                    if zipfile.is_zipfile(p):
                        with zipfile.ZipFile(p) as embedded: is_word='word/document.xml' in embedded.namelist()
                        if is_word: extract_docx(ctx,p,a['id'],b['id'],depth+1)
                        else: ctx.warn('unparsed_attachment',target)
                    else: ctx.warn('unparsed_attachment',target)
            def paragraph(node,owner):
                text=text_of(node); pr=node.find('w:pPr',NS); payload={}; kind='paragraph'
                if pr is not None:
                    sty=pr.find('w:pStyle',NS); style=styles.get(sty.get(W+'val') if sty is not None else '',{})
                    if sty is not None: payload['style_id']=sty.get(W+'val')
                    if style.get('level') or re.search(r'heading|заголовок',style.get('name',''),re.I):
                        kind='heading'; payload['level']=style.get('level') or 1
                    np=pr.find('w:numPr',NS)
                    if np is not None:
                        kind='list_item'; ni=np.find('w:numId',NS); lv=np.find('w:ilvl',NS)
                        num=ni.get(W+'val') if ni is not None else ''; level=int(lv.get(W+'val')) if lv is not None else 0
                        definition=nums.get(num,{}).get(level,{})
                        counters[num][level]+=1
                        payload['numbering']={'num_id':num,'level':level,'definition':definition,'sequence_count':counters[num][level]}
                fields=node.xpath('.//w:instrText/text()',namespaces=NS)
                if fields: payload['field_instructions']=fields
                hyperlinks=[]
                for link in node.findall('.//w:hyperlink',NS):
                    rid=link.get('{'+NS['r']+'}id'); rel=rels.get(rid)
                    if rel and rel['external']: hyperlinks.append(rel['target'])
                if hyperlinks: payload['hyperlinks']=hyperlinks
                b=ctx.block(kind,text,[xml_ref(ctx,aid,part,node)],owner,role,payload) if text or payload else None
                media(node,b['id'] if b else owner)
                for box in node.findall('.//w:txbxContent',NS): walk(box,b['id'] if b else owner)
            def table(node,owner):
                trs=node.findall('w:tr',NS); cells=[]; active={}
                b=ctx.block('table','',[xml_ref(ctx,aid,part,node)],owner,role,{'rows':len(trs),'columns':0,'cells':cells,'merges':[]})
                for row,tr in enumerate(trs):
                    col=0; next_active={}
                    before=tr.find('w:trPr/w:gridBefore',NS)
                    if before is not None: col=int(before.get(W+'val','0'))
                    for tc in tr.findall('w:tc',NS):
                        span=tc.find('w:tcPr/w:gridSpan',NS); width=int(span.get(W+'val','1')) if span is not None else 1
                        vm=tc.find('w:tcPr/w:vMerge',NS); continuation=vm is not None and vm.get(W+'val')!='restart'
                        paras=tc.findall('w:p',NS); text='\n'.join(text_of(p) for p in paras)
                        if continuation and col in active:
                            cell=active[col]; cell['rowspan']+=1; next_active[col]=cell
                            cell['source_refs'].append(xml_ref(ctx,aid,part,tc))
                            if text.strip():
                                cell['text']+='\n'+text; cell['value']=cell['text']; ctx.warn('merged_continuation_text','Text retained from a merged continuation cell.')
                        else:
                            cell={'id':b['id']+f':r{row}c{col}','row':row,'column':col,'rowspan':1,'colspan':width,'type':'string' if text else 'blank',
                                  'value':text if text else None,'text':text,'formula':None,'cached_value':None,'source_refs':[xml_ref(ctx,aid,part,tc)]}
                            cells.append(cell)
                            if vm is not None: next_active[col]=cell
                        for p in paras: media(p,b['id'])
                        for child in tc:
                            if child.tag==W+'tbl': table(child,b['id'])
                            elif child.tag!=W+'p' and child.tag!=W+'tcPr': walk(child,b['id'])
                        col+=width
                    active=next_active; b['payload']['columns']=max(b['payload']['columns'],col)
                b['payload']['merges']=[{'row':c['row'],'column':c['column'],'rowspan':c['rowspan'],'colspan':c['colspan']} for c in cells if c['rowspan']>1 or c['colspan']>1]
            def walk(node,owner):
                for child in node:
                    if child.tag==W+'p': paragraph(child,owner)
                    elif child.tag==W+'tbl': table(child,owner)
                    elif child.tag in {W+'footnote',W+'endnote'} and child.get(W+'type') in {'separator','continuationSeparator'}: continue
                    elif E.QName(child).localname not in {'sectPr','pPr','tcPr','trPr'}: walk(child,owner)
            walk(root,parent)
        ctx.doc['extraction']['coverage']['ooxml_parts']=parts

def excel_cell(ctx,b,row,col,value,text,typ,ref,formula=None,cached=None,rowspan=1,colspan=1):
    return {'id':b['id']+f':r{row}c{col}','row':row,'column':col,'rowspan':rowspan,'colspan':colspan,'type':typ,'value':value,'text':text,
            'formula':formula,'cached_value':cached,'source_refs':[ref]}

def extract_xlsx(ctx,path,aid='original'):
    from openpyxl.utils.cell import coordinate_to_tuple,range_boundaries
    with zipfile.ZipFile(native(path)) as z:
        strings=[]
        if 'xl/sharedStrings.xml' in z.namelist():
            for si in parse_xml(z.read('xl/sharedStrings.xml')): strings.append(''.join(si.xpath('.//s:t/text()',namespaces=NS)))
        book=parse_xml(z.read('xl/workbook.xml')); rels=relationships(z,'xl/workbook.xml'); sheets=[]
        for sheet in book.findall('s:sheets/s:sheet',NS):
            name=sheet.get('name'); target=rels[sheet.get('{'+NS['r']+'}id')]['target']; part=posixpath.normpath(posixpath.join('xl',target.lstrip('/'))) if not target.startswith('/') else target.lstrip('/')
            sheets.append(name); cells=[]
            b=ctx.block('table','',[ctx.loc(aid,'sheet_range',sheet=name,range='A1')],payload={'rows':0,'columns':0,'cells':cells,'merges':[],'sheet':name,'state':sheet.get('state','visible')})
            merges=[]; drawings=[]
            with z.open(part) as source:
                for _,node in E.iterparse(source,events=('end',),resolve_entities=False,no_network=True,huge_tree=True):
                    tag=E.QName(node).localname
                    if tag=='c':
                        address=node.get('r'); value=node.find('s:v',NS); inline=node.find('s:is',NS); formula=node.find('s:f',NS)
                        if value is not None or inline is not None or formula is not None:
                            row,col=coordinate_to_tuple(address); raw=value.text if value is not None else None; typ=node.get('t','n')
                            if typ=='s': val=strings[int(raw)] if raw is not None else None
                            elif typ=='inlineStr': val=''.join(inline.xpath('.//s:t/text()',namespaces=NS))
                            elif typ=='b': val=raw=='1' if raw is not None else None
                            elif typ in {'str','e','d'}: val=raw
                            else:
                                try: val=float(raw) if raw is not None else None; val=int(val) if val is not None and math.isfinite(val) and val.is_integer() else val
                                except (ValueError,TypeError): val=raw
                                if isinstance(val,float) and not math.isfinite(val): val=None; ctx.warn('nonfinite_number',name+'!'+address)
                            ftext='='+formula.text if formula is not None and formula.text else None
                            if formula is not None and not ftext: ctx.warn('shared_formula','Shared/array formula retains attributes at '+name+'!'+address)
                            cell=excel_cell(ctx,b,row-1,col-1,val,str(val) if val is not None else '',typ,ctx.loc(aid,'sheet_range',sheet=name,range=address),ftext,val if formula is not None else None)
                            cells.append(cell); b['payload']['rows']=max(b['payload']['rows'],row); b['payload']['columns']=max(b['payload']['columns'],col)
                            if formula is not None and raw is None: ctx.warn('formula_cache_missing',name+'!'+address)
                        node.clear()
                    elif tag=='mergeCell': merges.append(node.get('ref')); node.clear()
                    elif tag=='drawing': drawings.append(node.get('{'+NS['r']+'}id')); node.clear()
                    elif tag=='row': node.clear()
            b['payload']['merges']=merges
            for cell in cells:
                for merge in merges:
                    c1,r1,c2,r2=range_boundaries(merge)
                    if cell['row']==r1-1 and cell['column']==c1-1: cell['rowspan']=r2-r1+1; cell['colspan']=c2-c1+1
            if b['payload']['rows'] and b['payload']['columns']:
                from openpyxl.utils.cell import get_column_letter
                b['source_refs'][0]['range']='A1:'+get_column_letter(b['payload']['columns'])+str(b['payload']['rows'])
            sr=relationships(z,part)
            for rid in drawings:
                if rid not in sr: continue
                dp=posixpath.normpath(posixpath.join(posixpath.dirname(part),sr[rid]['target']))
                dr=selected(parse_xml(z.read(dp))); drel=relationships(z,dp); shapes=[]; connectors=[]; refs=[]
                for anchor in dr:
                    anch={E.QName(c).localname:{E.QName(v).localname:v.text for v in c} for c in anchor if E.QName(c).localname in {'from','to'}}
                    for obj in anchor:
                        typ=E.QName(obj).localname
                        if typ in {'sp','cxnSp','pic','graphicFrame','grpSp'}:
                            props=obj.xpath('.//*[local-name()="cNvPr"]'); oid=props[0].get('id') if props else None
                            ref=ctx.loc(aid,'drawing',part=dp,xpath=obj.getroottree().getpath(obj),object_id=oid,anchor=anch,sheet=name)
                            refs.append(ref); text='\n'.join(obj.xpath('.//a:p/a:r/a:t/text()',namespaces=NS))
                            if typ=='cxnSp':
                                ends=obj.xpath('.//a:stCxn|.//a:endCxn',namespaces=NS)
                                connectors.append({'id':oid,'endpoints':[dict(e.attrib) for e in ends],'anchor':anch,'source_refs':[ref]})
                            else: shapes.append({'id':oid,'kind':typ,'text':text,'anchor':anch,'geometry_xml':E.tostring(obj,encoding='unicode'),'source_refs':[ref]})
                            for blip in obj.findall('.//a:blip',NS):
                                link=drel.get(blip.get('{'+NS['r']+'}embed'))
                                if link and not link['external']:
                                    media=posixpath.normpath(posixpath.join(posixpath.dirname(dp),link['target']))
                                    a,p=ctx.save_bytes(z.read(media),Path(media).suffix.lstrip('.'),aid,'image',[ref])
                                    ctx.block('image','',[ref],b['id'],payload={'artifact_id':a['id'],'asset_id':'asset:'+a['sha256']})
                if refs: ctx.block('diagram','\n'.join(s['text'] for s in shapes if s['text']),refs,b['id'],payload={'shapes':shapes,'connectors':connectors,'sheet':name})
        ctx.doc['extraction']['coverage']['sheets']=sheets

def extract_xls(ctx,path):
    import xlrd
    book=xlrd.open_workbook(file_contents=Path(native(path)).read_bytes(),formatting_info=True,on_demand=True)
    sheets=[]
    for sheet in book.sheets():
        sheets.append(sheet.name); cells=[]
        b=ctx.block('table','',[ctx.loc(kind='sheet_range',sheet=sheet.name,range='A1')],payload={'rows':sheet.nrows,'columns':sheet.ncols,'cells':cells,'merges':[list(m) for m in sheet.merged_cells],'sheet':sheet.name})
        from openpyxl.utils.cell import get_column_letter
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                cell=sheet.cell(row,col)
                if cell.ctype in {xlrd.XL_CELL_EMPTY,xlrd.XL_CELL_BLANK}: continue
                val=cell.value
                if isinstance(val,float) and not math.isfinite(val): val=None
                typ={1:'string',2:'number',3:'excel_date_serial',4:'boolean',5:'error'}.get(cell.ctype,'unknown')
                span=next(((r2-r1,c2-c1) for r1,r2,c1,c2 in sheet.merged_cells if row==r1 and col==c1),(1,1))
                c=excel_cell(ctx,b,row,col,val,str(val),typ,ctx.loc(kind='sheet_range',sheet=sheet.name,range=get_column_letter(col+1)+str(row+1)),rowspan=span[0],colspan=span[1]); cells.append(c)
        b['payload']['date_system']=book.datemode
    ctx.doc['extraction']['coverage']['sheets']=sheets
    # xlrd exposes cached results but not BIFF formulas or drawings. Preserve a derivative view and flag this limitation.
    ctx.warn('xls_native_limitations','BIFF formula expressions/drawing geometry require the preserved XLSX conversion; original cell references remain authoritative.')

def office_convert(ctx,source,work,tools,ext):
    global _OFFICE_STATE
    import shutil
    work=Path(work); work.mkdir(parents=True,exist_ok=True)
    staged=work/('input'+Path(source).suffix.lower()); shutil.copyfile(native(source),staged)
    outfile=work/('input.'+ext)
    profile=work.parent/('office_worker_'+str(os.getpid()))
    (profile/'user').mkdir(parents=True,exist_ok=True)
    # Preserve cached source values; do not update Writer fields or Calc formulas on load.
    env=dict(os.environ,SAL_USE_VCLPLUGIN='svp',SATBAYEV_OFFICE_KEEPALIVE='1')
    if _OFFICE_STATE is None or _OFFICE_STATE['listener'].poll() is not None:
        (profile/'user'/'registrymodifications.xcu').write_text('''<?xml version="1.0" encoding="UTF-8"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Field" oor:op="fuse"><value>false</value></prop><prop oor:name="Chart" oor:op="fuse"><value>false</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Formula"><prop oor:name="ExcelRecalcMode" oor:op="fuse"><value>1</value></prop></item></oor:items>''',encoding='utf-8')
        pipe='satbayev_'+uuid.uuid4().hex
        args=[tools['libreoffice'],'-env:UserInstallation='+profile.as_uri(),'--headless','--nologo','--nodefault','--nolockcheck','--norestore','--accept=pipe,name='+pipe+';urp;StarOffice.ComponentContext']
        listener=subprocess.Popen(args,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,env=env,creationflags=CREATE_NO_WINDOW)
        _OFFICE_STATE={'listener':listener,'pipe':pipe}
        from multiprocessing.util import Finalize
        Finalize(None,stop_office,exitpriority=10)
    bridge=Path(__file__).with_name('office_bridge.py'); py=Path(tools['libreoffice']).parent/'python.exe'
    try:
        p=subprocess.run([str(py),str(bridge),_OFFICE_STATE['pipe'],str(staged),str(outfile),ext],capture_output=True,timeout=180,env=env,creationflags=CREATE_NO_WINDOW)
        if p.returncode or not outfile.exists():raise RuntimeError('LibreOffice conversion failed: '+p.stderr.decode('utf-8','replace')[-1000:])
    except Exception:
        stop_office();raise
    a,path=ctx.save_bytes(outfile.read_bytes(),ext,'original','office_conversion')
    return a,path
